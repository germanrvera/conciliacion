"""
excel_export.py — Genera el Excel de resultado de 7 solapas
"""

from datetime import datetime
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side

NUM = '#,##0.00;[Red]-#,##0.00'
DF = 'DD/MM/YYYY'
T = '1F3864'; H2 = '2F5496'; H3 = 'BDD7EE'
EX = 'C6EFCE'; FE = 'DDEBF7'; MU = 'FFCC99'; PE = 'FFE0E0'; GA = 'FCE4D6'
DP_C = 'D9EAD3'; MS_C = 'CFE2F3'; REV_C = 'EAD1DC'; MANUAL_C = 'FFE599'
A1 = 'F5F5F5'; A2 = 'FFFFFF'; OK = '375623'; WA = '9C0006'


def _bg_de(nivel):
    return {
        'EXACTO': EX, 'FECHA': FE, 'MULTIPLES': MU, 'DP': DP_C,
        'MULTI_SAP': MS_C, 'REVERSION': REV_C, 'CANCELACION': REV_C,
        'MANUAL': MANUAL_C, 'SIN_CRUCE_CONFIRMADO': A1,
    }.get(nivel, A1)


def _F(h): return PatternFill("solid", fgColor=h)
def _Ft(bold=False, color="000000", size=10, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic, name="Arial")
def _Al(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)
def _Bd(c="D0D0D0"):
    s = Side(style='thin', color=c)
    return Border(left=s, right=s, top=s, bottom=s)


def _W(ws, r, c, val, bg=None, fnt=None, aln=None, fmt=None, brd=None, span=None):
    if span and span > 1:
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + span - 1)
        for ci in range(c + 1, c + span):
            mc = ws.cell(row=r, column=ci)
            if bg: mc.fill = _F(bg)
    cell = ws.cell(row=r, column=c, value=val)
    if bg: cell.fill = _F(bg)
    if fnt: cell.font = fnt
    if aln: cell.alignment = aln
    if fmt: cell.number_format = fmt
    if brd: cell.border = brd
    return cell


def generar_excel(resultado, banco_nombre='Credicoop', out_path=None):
    """
    resultado: el dict que devuelve motor.conciliar()
    out_path: si se pasa, guarda ahí. Si no, devuelve el Workbook.
    """
    mayor = resultado['mayor']; extracto = resultado['extracto']
    cruces = resultado['cruces']; cruces_bco = resultado['cruces_bco']
    sin_sap = resultado['sin_sap']; sin_bco = resultado['sin_bco']
    gastos = resultado['gastos']; niveles = resultado['niveles']
    SALDO_SAP = resultado['SALDO_SAP']; SALDO_BCO = resultado['SALDO_BCO']
    db_nc = resultado['db_nc']; cr_nc = resultado['cr_nc']
    db_c = resultado['db_c']; cr_c = resultado['cr_c']
    sc_bco = resultado['sc_bco']; sc_sap = resultado['sc_sap']; dif = resultado['dif']
    sap_real = mayor[1:] if mayor else []

    wb = Workbook()

    # ── RESUMEN ────────────────────────────────────────────────────────
    ws = wb.active; ws.title = "RESUMEN"
    ws.sheet_properties.tabColor = T; ws.sheet_view.showGridLines = False
    for col, w in [('A', 2), ('B', 30), ('C', 20), ('D', 20), ('E', 4),
                   ('F', 20), ('G', 20), ('H', 20), ('I', 2)]:
        ws.column_dimensions[col].width = w
    ws.row_dimensions[1].height = 6; ws.row_dimensions[2].height = 22
    _W(ws, 2, 2, f'CONCILIACIÓN BANCARIA — {banco_nombre.upper()}',
       bg=T, fnt=_Ft(True, 'FFFFFF', 13), aln=_Al("left"), span=6)
    _W(ws, 2, 8, datetime.now().strftime('%d/%m/%Y'), bg=T, fnt=_Ft(False, 'AAAAAA', 9), aln=_Al("right"))

    ws.row_dimensions[3].height = 8; ws.row_dimensions[4].height = 26
    for c, l in [(2, ''), (3, 'SALDO INICIAL'), (4, 'DB NO CONTAB.'), (6, 'CR NO CONTAB.'),
                 (7, 'DB CONTABLES'), (8, 'CR CONTABLES')]:
        _W(ws, 4, c, l, bg=H2, fnt=_Ft(True, 'FFFFFF', 9), aln=_Al(wrap=True))
    _W(ws, 4, 5, '', bg=H2)

    ws.row_dimensions[5].height = 15
    _W(ws, 5, 2, 'BANCO', bg=A1, fnt=_Ft(True), aln=_Al("left"))
    _W(ws, 5, 3, SALDO_BCO, bg=A1, fnt=_Ft(), aln=_Al("right"), fmt=NUM)
    _W(ws, 5, 4, '', bg=A1); _W(ws, 5, 5, '', bg=A1); _W(ws, 5, 6, '', bg=A1)
    _W(ws, 5, 7, db_c, bg='FFF2CC', fnt=_Ft(), aln=_Al("right"), fmt=NUM)
    _W(ws, 5, 8, cr_c, bg='FFF2CC', fnt=_Ft(), aln=_Al("right"), fmt=NUM)

    ws.row_dimensions[6].height = 15
    _W(ws, 6, 2, 'CONTABILIDAD', bg=A2, fnt=_Ft(True), aln=_Al("left"))
    _W(ws, 6, 3, SALDO_SAP, bg=A2, fnt=_Ft(), aln=_Al("right"), fmt=NUM)
    _W(ws, 6, 4, db_nc, bg='FFF2CC', fnt=_Ft(), aln=_Al("right"), fmt=NUM)
    _W(ws, 6, 5, '', bg=A2)
    _W(ws, 6, 6, cr_nc, bg='FFF2CC', fnt=_Ft(), aln=_Al("right"), fmt=NUM)
    _W(ws, 6, 7, '', bg=A2); _W(ws, 6, 8, '', bg=A2)

    ws.row_dimensions[7].height = 15
    dc = PE if abs(dif) > 1 else EX
    dfnt = _Ft(True, WA if abs(dif) > 1 else OK)
    _W(ws, 7, 2, 'DIFERENCIAS', bg=dc, fnt=dfnt, aln=_Al("left"))
    _W(ws, 7, 3, SALDO_BCO - SALDO_SAP, bg=dc, fnt=dfnt, aln=_Al("right"), fmt=NUM)
    _W(ws, 7, 4, db_nc, bg=dc, fnt=dfnt, aln=_Al("right"), fmt=NUM); _W(ws, 7, 5, '', bg=dc)
    _W(ws, 7, 6, cr_nc, bg=dc, fnt=dfnt, aln=_Al("right"), fmt=NUM)
    _W(ws, 7, 7, db_c, bg=dc, fnt=dfnt, aln=_Al("right"), fmt=NUM)
    _W(ws, 7, 8, cr_c, bg=dc, fnt=dfnt, aln=_Al("right"), fmt=NUM)

    ws.row_dimensions[8].height = 8
    ws.row_dimensions[9].height = 15
    _W(ws, 9, 2, 'Saldo conciliado BANCO', bg=A1, fnt=_Ft(True, size=9), aln=_Al("left"))
    _W(ws, 9, 3, sc_bco, bg=FE, fnt=_Ft(True), aln=_Al("right"), fmt=NUM)
    ws.row_dimensions[10].height = 15
    _W(ws, 10, 2, 'Saldo conciliado SAP', bg=A1, fnt=_Ft(True, size=9), aln=_Al("left"))
    _W(ws, 10, 3, sc_sap, bg=FE, fnt=_Ft(True), aln=_Al("right"), fmt=NUM)
    ws.row_dimensions[11].height = 18
    dif_bg = EX if abs(dif) < 1 else PE
    dif_f = _Ft(True, OK if abs(dif) < 1 else WA, 12)
    _W(ws, 11, 2, 'DIFERENCIA CONCILIADA', bg=dif_bg, fnt=dif_f, aln=_Al("left"))
    _W(ws, 11, 3, dif, bg=dif_bg, fnt=dif_f, aln=_Al("right"), fmt=NUM)

    ws.row_dimensions[12].height = 8; ws.row_dimensions[13].height = 15
    _W(ws, 13, 2, 'Resultado del cruce automático', bg=H2, fnt=_Ft(True, 'FFFFFF', 9), aln=_Al("left"), span=2)
    filas_stats = [
        ('EXACTO', niveles.get('EXACTO', 0), EX),
        ('FECHA', niveles.get('FECHA', 0), FE),
        ('DP (1 SAP → varios banco)', niveles.get('DP', 0), DP_C),
        ('MULTI-SAP (varios SAP → 1 banco)', niveles.get('MULTI_SAP', 0), MS_C),
        ('Cruce manual (operador)', niveles.get('MANUAL', 0), MANUAL_C),
        ('Reversión / cancelación contable',
         niveles.get('REVERSION', 0) + niveles.get('CANCELACION', 0), REV_C),
        ('MÚLTIPLES — revisar', niveles.get('MULTIPLES', 0), MU),
        ('SAP sin cruce', len(sin_sap), PE),
        ('Banco sin cruce', len(sin_bco), PE),
        ('Gastos bancarios', len(gastos), GA),
    ]
    ri = 14
    for lbl, val, bg in filas_stats:
        ws.row_dimensions[ri].height = 13
        _W(ws, ri, 2, lbl, bg=bg, fnt=_Ft(size=9), aln=_Al("left"))
        _W(ws, ri, 3, val, bg=bg, fnt=_Ft(True, size=9), aln=_Al("center"), fmt='#,##0')
        ri += 1

    ri += 1
    leyenda = [
        (EX, 'EXACTO', 'Importe + fecha + CUIT, descripción o HAB'),
        (FE, 'FECHA', 'Importe + fecha dentro de tolerancia'),
        (DP_C, 'DP', 'Depósito cruzado con varias líneas banco'),
        (MS_C, 'MULTI-SAP', 'Varias partidas SAP suman 1 movimiento banco'),
        (MANUAL_C, 'MANUAL', 'Cruce confirmado manualmente por el operador'),
        (REV_C, 'REVERSIÓN/CANCEL.', 'Par contable interno — sin banco propio'),
        (MU, 'MÚLTIPLES', 'Varios candidatos o importe ,000 sin CUIT'),
        (PE, 'SIN CRUCE', 'Sin contraparte encontrada'),
        (GA, 'GASTO BANCARIO', 'Cargo bancario agrupado en PP de gastos'),
    ]
    for bg_l, lbl_l, dsc_l in leyenda:
        ws.row_dimensions[ri].height = 13
        _W(ws, ri, 2, lbl_l, bg=bg_l, fnt=_Ft(True, size=9), aln=_Al("center"))
        _W(ws, ri, 3, dsc_l, fnt=_Ft(size=9), aln=_Al("left"), span=5)
        ri += 1

    # ── PENDIENTES ─────────────────────────────────────────────────────
    wp = wb.create_sheet("PENDIENTES"); wp.sheet_properties.tabColor = 'FF0000'
    wp.sheet_view.showGridLines = False; wp.freeze_panes = "A3"
    for col, w in [('A', 12), ('B', 12), ('C', 52), ('D', 18)]:
        wp.column_dimensions[col].width = w
    wp.row_dimensions[1].height = 20
    _W(wp, 1, 1, '⚠  PARTIDAS PENDIENTES DE CONCILIACIÓN',
       bg='9C0006', fnt=_Ft(True, 'FFFFFF', 12), aln=_Al("left"), span=4)
    nr = 2

    multiples = [(s, cruces[s['idx']]) for s in sap_real
                 if s['idx'] in cruces and cruces[s['idx']]['nivel'] == 'MULTIPLES']
    if multiples:
        _W(wp, nr, 1, f'MÚLTIPLES — {len(multiples)} casos', bg=MU, fnt=_Ft(True, '7B3F00', 10),
           aln=_Al("left"), span=4); nr += 1
        wp.row_dimensions[nr].height = 20
        for c, l in [(1, 'NRO SAP'), (2, 'FECHA'), (3, 'DESCRIPCIÓN'), (4, 'IMPORTE')]:
            _W(wp, nr, c, l, bg=H3, fnt=_Ft(True, T, 9), aln=_Al())
        nr += 1
        for s, info in multiples:
            wp.row_dimensions[nr].height = 13
            _W(wp, nr, 1, s['ndoc'], bg=MU, fnt=_Ft(size=9), aln=_Al("center"), brd=_Bd())
            _W(wp, nr, 2, s['fecha'], bg=MU, fnt=_Ft(size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
            _W(wp, nr, 3, s['nombre'] or s['comentario'], bg=MU, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
            _W(wp, nr, 4, s['importe'], bg=MU, fnt=_Ft(size=9), aln=_Al("right"), fmt=NUM, brd=_Bd())
            nr += 1
        nr += 1

    _W(wp, nr, 1, f'SAP SIN CRUCE — {len(sin_sap)} partidas', bg=H2, fnt=_Ft(True, 'FFFFFF', 10),
       aln=_Al("left"), span=4); nr += 1
    wp.row_dimensions[nr].height = 20
    for c, l in [(1, 'NRO SAP'), (2, 'FECHA'), (3, 'DESCRIPCIÓN'), (4, 'IMPORTE')]:
        _W(wp, nr, c, l, bg=H3, fnt=_Ft(True, T, 9), aln=_Al())
    nr += 1
    for s in sin_sap:
        wp.row_dimensions[nr].height = 13
        bg = A1 if nr % 2 == 0 else A2
        _W(wp, nr, 1, s['ndoc'], bg=PE, fnt=_Ft(True, WA, size=9), aln=_Al("center"), brd=_Bd())
        _W(wp, nr, 2, s['fecha'], bg=bg, fnt=_Ft(size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
        _W(wp, nr, 3, s['nombre'] or s['comentario'], bg=bg, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
        _W(wp, nr, 4, s['importe'], bg=bg, fnt=_Ft(size=9), aln=_Al("right"), fmt=NUM, brd=_Bd())
        nr += 1
    nr += 1
    _W(wp, nr, 1, f'BANCO SIN CRUCE — {len(sin_bco)} partidas', bg=H2, fnt=_Ft(True, 'FFFFFF', 10),
       aln=_Al("left"), span=4); nr += 1
    wp.row_dimensions[nr].height = 20
    for c, l in [(1, 'FECHA'), (2, 'IMPORTE'), (3, 'CONCEPTO BANCO'), (4, 'TIPO')]:
        _W(wp, nr, c, l, bg=H3, fnt=_Ft(True, T, 9), aln=_Al())
    nr += 1
    for e in sin_bco:
        wp.row_dimensions[nr].height = 13
        bg = A1 if nr % 2 == 0 else A2
        _W(wp, nr, 1, e['fecha'], bg=PE, fnt=_Ft(True, WA, size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
        _W(wp, nr, 2, e['importe'], bg=bg, fnt=_Ft(size=9), aln=_Al("right"), fmt=NUM, brd=_Bd())
        _W(wp, nr, 3, e['concepto'], bg=bg, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
        _W(wp, nr, 4, 'DÉBITO' if e['importe'] < 0 else 'CRÉDITO', bg=bg, fnt=_Ft(size=9),
           aln=_Al("center"), brd=_Bd())
        nr += 1

    # ── CRUCES ─────────────────────────────────────────────────────────
    wc = wb.create_sheet("CRUCES"); wc.sheet_properties.tabColor = '00B050'
    wc.sheet_view.showGridLines = False; wc.freeze_panes = "A3"
    for col, w in [('A', 12), ('B', 12), ('C', 38), ('D', 16), ('E', 12),
                   ('F', 42), ('G', 7), ('H', 14), ('I', 42)]:
        wc.column_dimensions[col].width = w
    wc.row_dimensions[1].height = 18
    _W(wc, 1, 1, 'CRUCES AUTOMÁTICOS', bg=T, fnt=_Ft(True, 'FFFFFF', 11), aln=_Al("left"), span=9)
    wc.row_dimensions[2].height = 22
    for c, l in [(1, 'NRO SAP'), (2, 'FECHA SAP'), (3, 'DESC SAP'), (4, 'IMPORTE'),
                 (5, 'FECHA BCO'), (6, 'CONCEPTO BANCO'), (7, 'DÍAS'), (8, 'NIVEL'), (9, 'MOTIVO')]:
        _W(wc, 2, c, l, bg=H3, fnt=_Ft(True, T, 9), aln=_Al(wrap=True))
    row = 3
    for s in sap_real:
        if s['idx'] not in cruces:
            continue
        info = cruces[s['idx']]; nv = info['nivel']
        if nv in ('REVERSION', 'CANCELACION', 'SIN_CRUCE_CONFIRMADO'):
            continue
        bg = _bg_de(nv)
        wc.row_dimensions[row].height = 13
        if info.get('bco_list') and len(info['bco_list']) > 1:
            conc_bco = f"[{len(info['bco_list'])} líneas] " + info['bco_list'][0]['concepto'][:32]
        elif info.get('bco'):
            conc_bco = info['bco']['concepto']
        else:
            conc_bco = ''
        _W(wc, row, 1, s['ndoc'], bg=bg, fnt=_Ft(size=9), aln=_Al("center"), brd=_Bd())
        _W(wc, row, 2, s['fecha'], bg=bg, fnt=_Ft(size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
        _W(wc, row, 3, s['nombre'] or s['comentario'], bg=bg, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
        _W(wc, row, 4, s['importe'], bg=bg, fnt=_Ft(size=9), aln=_Al("right"), fmt=NUM, brd=_Bd())
        _W(wc, row, 5, info['bco']['fecha'] if info.get('bco') else None,
           bg=bg, fnt=_Ft(size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
        _W(wc, row, 6, conc_bco, bg=bg, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
        _W(wc, row, 7, info['dias'], bg=bg, fnt=_Ft(size=9), aln=_Al("center"), fmt='#,##0', brd=_Bd())
        _W(wc, row, 8, nv, bg=bg, fnt=_Ft(True, size=9), aln=_Al("center"), brd=_Bd())
        _W(wc, row, 9, info['motivo'], bg=bg, fnt=_Ft(size=8, italic=True), aln=_Al("left", wrap=True), brd=_Bd())
        row += 1

    # ── REVERSIONES ────────────────────────────────────────────────────
    wr = wb.create_sheet("REVERSIONES"); wr.sheet_properties.tabColor = REV_C
    wr.sheet_view.showGridLines = False; wr.freeze_panes = "A3"
    for col, w in [('A', 12), ('B', 12), ('C', 45), ('D', 18), ('E', 35)]:
        wr.column_dimensions[col].width = w
    wr.row_dimensions[1].height = 18
    _W(wr, 1, 1, 'REVERSIONES Y CANCELACIONES CONTABLES — sin movimiento bancario propio',
       bg='8E7CC3', fnt=_Ft(True, 'FFFFFF', 11), aln=_Al("left"), span=5)
    wr.row_dimensions[2].height = 20
    for c, l in [(1, 'NRO SAP'), (2, 'FECHA'), (3, 'DESCRIPCIÓN'), (4, 'IMPORTE'), (5, 'MOTIVO')]:
        _W(wr, 2, c, l, bg=H3, fnt=_Ft(True, T, 9), aln=_Al(wrap=True))
    row = 3
    for s in sap_real:
        if s['idx'] not in cruces:
            continue
        info = cruces[s['idx']]
        if info['nivel'] not in ('REVERSION', 'CANCELACION'):
            continue
        wr.row_dimensions[row].height = 13
        _W(wr, row, 1, s['ndoc'], bg=REV_C, fnt=_Ft(size=9), aln=_Al("center"), brd=_Bd())
        _W(wr, row, 2, s['fecha'], bg=REV_C, fnt=_Ft(size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
        _W(wr, row, 3, s['nombre'] or s['comentario'], bg=REV_C, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
        _W(wr, row, 4, s['importe'], bg=REV_C, fnt=_Ft(size=9), aln=_Al("right"), fmt=NUM, brd=_Bd())
        _W(wr, row, 5, info['motivo'], bg=REV_C, fnt=_Ft(size=8, italic=True), aln=_Al("left", wrap=True), brd=_Bd())
        row += 1

    # ── GASTOS BANCARIOS ───────────────────────────────────────────────
    wg = wb.create_sheet("GASTOS BANCARIOS"); wg.sheet_properties.tabColor = 'A9573B'
    wg.sheet_view.showGridLines = False; wg.freeze_panes = "A3"
    for col, w in [('A', 12), ('B', 55), ('C', 18), ('D', 22)]:
        wg.column_dimensions[col].width = w
    wg.row_dimensions[1].height = 18
    _W(wg, 1, 1, 'GASTOS BANCARIOS', bg='A9573B', fnt=_Ft(True, 'FFFFFF', 11), aln=_Al("left"), span=4)
    wg.row_dimensions[2].height = 20
    for c, l in [(1, 'FECHA'), (2, 'CONCEPTO'), (3, 'IMPORTE'), (4, 'TIPO')]:
        _W(wg, 2, c, l, bg=H3, fnt=_Ft(True, T, 9), aln=_Al())
    row = 3; tot_g = defaultdict(float)
    for e in gastos:
        wg.row_dimensions[row].height = 13
        bg = A1 if row % 2 == 0 else A2
        _W(wg, row, 1, e['fecha'], bg=bg, fnt=_Ft(size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
        _W(wg, row, 2, e['concepto'], bg=bg, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
        _W(wg, row, 3, e['importe'], bg=bg, fnt=_Ft(size=9), aln=_Al("right"), fmt=NUM, brd=_Bd())
        _W(wg, row, 4, e['gasto'], bg=GA, fnt=_Ft(size=8, italic=True), aln=_Al("left"), brd=_Bd())
        tot_g[e['gasto']] += e['importe']; row += 1
    row += 1
    _W(wg, row, 1, 'TOTALES', bg=H2, fnt=_Ft(True, 'FFFFFF', 9), aln=_Al(), span=4); row += 1
    for tipo, tot in sorted(tot_g.items(), key=lambda x: x[1]):
        wg.row_dimensions[row].height = 13
        _W(wg, row, 2, tipo, bg=GA, fnt=_Ft(size=9), aln=_Al("left"))
        _W(wg, row, 3, tot, bg=GA, fnt=_Ft(True, size=9), aln=_Al("right"), fmt=NUM)
        row += 1

    # ── MAYOR ──────────────────────────────────────────────────────────
    wm = wb.create_sheet("MAYOR"); wm.sheet_properties.tabColor = H2
    wm.sheet_view.showGridLines = False; wm.freeze_panes = "A3"
    for col, w in [('A', 12), ('B', 12), ('C', 45), ('D', 18), ('E', 14), ('F', 40)]:
        wm.column_dimensions[col].width = w
    wm.row_dimensions[1].height = 18
    _W(wm, 1, 1, 'MAYOR SAP', bg=H2, fnt=_Ft(True, 'FFFFFF', 11), aln=_Al("left"), span=6)
    wm.row_dimensions[2].height = 20
    for c, l in [(1, 'NRO DOC'), (2, 'FECHA'), (3, 'NOMBRE / DESCRIPCIÓN'), (4, 'IMPORTE'),
                 (5, 'ESTADO'), (6, 'CRUCE CON')]:
        _W(wm, 2, c, l, bg=H3, fnt=_Ft(True, T, 9), aln=_Al(wrap=True))
    row = 3
    for s in sap_real:
        wm.row_dimensions[row].height = 13
        info = cruces.get(s['idx'])
        if info:
            nv = info['nivel']; bg = _bg_de(nv); cruce = info['motivo'][:45]; estado = nv
        else:
            bg = PE; cruce = ''; estado = 'SIN CRUCE'
        _W(wm, row, 1, s['ndoc'], bg=bg, fnt=_Ft(size=9), aln=_Al("center"), brd=_Bd())
        _W(wm, row, 2, s['fecha'], bg=bg, fnt=_Ft(size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
        _W(wm, row, 3, s['nombre'] or s['comentario'], bg=bg, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
        _W(wm, row, 4, s['importe'], bg=bg, fnt=_Ft(size=9), aln=_Al("right"), fmt=NUM, brd=_Bd())
        _W(wm, row, 5, estado, bg=bg, fnt=_Ft(bold=estado in ('SIN CRUCE', 'MULTIPLES'), size=9),
           aln=_Al("center"), brd=_Bd())
        _W(wm, row, 6, cruce, bg=bg, fnt=_Ft(size=8, italic=True), aln=_Al("left"), brd=_Bd())
        row += 1

    # ── EXTRACTO ───────────────────────────────────────────────────────
    we = wb.create_sheet("EXTRACTO"); we.sheet_properties.tabColor = '375623'
    we.sheet_view.showGridLines = False; we.freeze_panes = "A3"
    for col, w in [('A', 12), ('B', 55), ('C', 18), ('D', 14), ('E', 25)]:
        we.column_dimensions[col].width = w
    we.row_dimensions[1].height = 18
    _W(we, 1, 1, f'EXTRACTO BANCARIO — {banco_nombre.upper()}', bg='375623',
       fnt=_Ft(True, 'FFFFFF', 11), aln=_Al("left"), span=5)
    we.row_dimensions[2].height = 20
    for c, l in [(1, 'FECHA'), (2, 'CONCEPTO'), (3, 'IMPORTE'), (4, 'ESTADO'), (5, 'CRUCE CON')]:
        _W(we, 2, c, l, bg=H3, fnt=_Ft(True, T, 9), aln=_Al(wrap=True))
    row = 3
    for e in extracto:
        we.row_dimensions[row].height = 13
        if e['gasto']:
            bg = GA; estado = 'GASTO'; cruce = e['gasto']
        elif e['idx'] in cruces_bco:
            idx_s = cruces_bco[e['idx']]; nv = cruces[idx_s]['nivel']
            bg = _bg_de(nv); estado = nv
            sr = next((s for s in mayor if s['idx'] == idx_s), None)
            cruce = sr['ndoc'] if sr else ''
        else:
            bg = PE; estado = 'SIN CRUCE'; cruce = ''
        _W(we, row, 1, e['fecha'], bg=bg, fnt=_Ft(size=9), aln=_Al("center"), fmt=DF, brd=_Bd())
        _W(we, row, 2, e['concepto'], bg=bg, fnt=_Ft(size=8), aln=_Al("left"), brd=_Bd())
        _W(we, row, 3, e['importe'], bg=bg, fnt=_Ft(size=9), aln=_Al("right"), fmt=NUM, brd=_Bd())
        _W(we, row, 4, estado, bg=bg, fnt=_Ft(bold=estado == 'SIN CRUCE', size=9), aln=_Al("center"), brd=_Bd())
        _W(we, row, 5, cruce, bg=bg, fnt=_Ft(size=8, italic=True), aln=_Al("left"), brd=_Bd())
        row += 1

    if out_path:
        wb.save(out_path)
        return out_path
    return wb
